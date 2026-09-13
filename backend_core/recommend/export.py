"""推荐简报 Excel / PDF 导出。"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any, Dict, List, Optional


def _score_detail_text(detail: Any) -> str:
    if not isinstance(detail, dict):
        return ""
    parts = []
    for key in ("resonance_note", "quality_note", "action_note", "role_note", "note"):
        v = detail.get(key)
        if not v:
            continue
        if key == "resonance_note":
            parts.append(f"{v}={detail.get('resonance', '')}")
        elif key == "quality_note":
            parts.append(f"{v}={detail.get('quality', '')}")
        elif key == "role_note":
            parts.append(f"{v}={detail.get('role_bonus', '')}")
        else:
            parts.append(str(v))
    total = detail.get("total")
    if total is not None:
        parts.append(f"合计={total}")
    return "；".join(parts)


def _zone_text(zone: Any) -> str:
    if not isinstance(zone, dict):
        return ""
    parts = []
    for k in ("low", "price", "high"):
        if zone.get(k) is not None:
            parts.append(f"{k}={zone.get(k)}")
    label = zone.get("label") or ""
    return (label + " " + ",".join(parts)).strip()


def build_recommend_excel(
    brief: Dict[str, Any],
    *,
    output_dir: Optional[str] = None,
) -> str:
    """生成 Excel，返回文件路径。"""
    import pandas as pd

    horizon = str(brief.get("horizon") or "daily")
    asof = str(brief.get("asof_date") or datetime.now().strftime("%Y-%m-%d"))
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = output_dir or os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
        "exported_docs",
    )
    os.makedirs(out_dir, exist_ok=True)
    filename = f"recommend_{horizon}_{asof}_{ts}.xlsx"
    filepath = os.path.join(out_dir, filename)

    items = list(brief.get("items") or [])
    exec_rows = []
    watch_rows = []
    for it in items:
        row = {
            "代码": it.get("code"),
            "名称": it.get("name"),
            "立场": it.get("stance") or it.get("action"),
            "主策略": it.get("primary_strategy"),
            "策略共振": ",".join(it.get("strategies") or []),
            "角色": it.get("role_label") or it.get("role"),
            "行业": it.get("industry") or it.get("board_name") or it.get("board_code"),
            "推荐分": it.get("recommend_score"),
            "推荐分明细": _score_detail_text(it.get("score_detail")),
            "买区": _zone_text(it.get("buy_zone")),
            "止损": _zone_text(it.get("stop_zone")),
            "止盈": _zone_text(it.get("take_profit")),
            "仓位建议": it.get("position_hint"),
            "摘要": it.get("summary"),
            "约束": ",".join(it.get("constraint_reasons") or []),
        }
        if it.get("action") == "buy":
            exec_rows.append(row)
        else:
            watch_rows.append(row)

    summary = brief.get("summary") or {}
    summary_rows = [
        {"项": "horizon", "值": horizon},
        {"项": "asof_date", "值": asof},
        {"项": "plan_for", "值": brief.get("plan_for")},
        {"项": "market_stance", "值": brief.get("market_stance")},
        {"项": "disclaimer", "值": summary.get("disclaimer")},
        {"项": "executable", "值": len(exec_rows)},
        {"项": "watch", "值": len(watch_rows)},
    ]

    risk_rows = []
    for r in brief.get("risk_observe") or []:
        risk_rows.append(
            {
                "代码": r.get("code"),
                "名称": r.get("name"),
                "来源": r.get("source"),
                "类型": r.get("kind"),
                "备注": r.get("note"),
            }
        )

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        pd.DataFrame(summary_rows).to_excel(writer, index=False, sheet_name="摘要")
        pd.DataFrame(exec_rows or [{"提示": "无可执行"}]).to_excel(
            writer, index=False, sheet_name="可执行"
        )
        pd.DataFrame(watch_rows or [{"提示": "无观察"}]).to_excel(
            writer, index=False, sheet_name="观察"
        )
        if risk_rows:
            pd.DataFrame(risk_rows).to_excel(writer, index=False, sheet_name="结构风险观察")
        kpi = brief.get("kpi") or {}
        if kpi:
            pd.DataFrame([{"键": k, "值": str(v)} for k, v in kpi.items()]).to_excel(
                writer, index=False, sheet_name="KPI"
            )

    # 简单列宽
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        wb = load_workbook(filepath)
        for ws in wb.worksheets:
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                maxlen = 10
                for cell in col[:50]:
                    maxlen = max(maxlen, min(len(str(cell.value or "")), 40))
                ws.column_dimensions[letter].width = maxlen + 2
        wb.save(filepath)
    except Exception:
        pass
    return filepath


def build_recommend_pdf_bytes(brief: Dict[str, Any]) -> bytes:
    """服务端简易 PDF（reportlab）；不可用则抛错由调用方回退。"""
    from io import BytesIO

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    # 尝试注册中文字体
    font_name = "Helvetica"
    for fp in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ):
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("RecommendCN", fp))
                font_name = "RecommendCN"
                break
            except Exception:
                continue

    y = height - 20 * mm
    horizon = brief.get("horizon")
    asof = brief.get("asof_date")
    c.setFont(font_name, 14)
    c.drawString(20 * mm, y, f"策略推荐简报 ({horizon})  asof={asof}")
    y -= 8 * mm
    c.setFont(font_name, 9)
    c.drawString(
        20 * mm,
        y,
        f"立场环境: {brief.get('market_stance') or '-'}  plan_for={brief.get('plan_for') or '-'}",
    )
    y -= 6 * mm
    c.drawString(20 * mm, y, "规则合成参考，非投资建议")
    y -= 10 * mm

    c.setFont(font_name, 10)
    c.drawString(20 * mm, y, "可执行 / 观察清单")
    y -= 6 * mm
    c.setFont(font_name, 8)
    for it in (brief.get("items") or [])[:40]:
        line = (
            f"{it.get('code')} {it.get('name') or ''} | {it.get('stance')} | "
            f"{it.get('role_label') or it.get('role')} | {it.get('primary_strategy') or '-'} | "
            f"分={it.get('recommend_score')}"
        )
        if y < 20 * mm:
            c.showPage()
            y = height - 20 * mm
            c.setFont(font_name, 8)
        c.drawString(20 * mm, y, line[:110])
        y -= 5 * mm

    c.showPage()
    c.save()
    return buf.getvalue()
